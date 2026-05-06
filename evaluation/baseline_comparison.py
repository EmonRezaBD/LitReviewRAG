"""Cost-quality baseline comparison: GPT-4o-mini vs Llama-3.1-70b-instruct.

Runs the full extraction pipeline twice over the 5 evaluation papers —
once with OpenAI's GPT-4o-mini and once with Meta's Llama-3.1-70b-instruct
via OpenRouter. For each run it records:

- BERTScore F1 against the human-annotated ground truth.
- Total wall-clock latency.
- Approximate API cost (estimated from per-million-token pricing).

The output is a side-by-side comparison table appended to
docs/evaluation_results.md, characterizing the cost/quality trade-off
between the two models for the structured-extraction task.

Usage:
    python -m evaluation.baseline_comparison
"""

from __future__ import annotations

import json
import logging
import time
from dataclasses import dataclass
from pathlib import Path

from bert_score import score as bert_score
from openpyxl import load_workbook

from litreviewrag import config
from litreviewrag.extraction.extractor import extract_paper
from litreviewrag.extraction.prompts import FIELD_SPECS
from litreviewrag.retrieval.hybrid_search import HybridSearcher
from litreviewrag.retrieval.vector_store import VectorStore

logging.basicConfig(level=logging.WARNING, format="%(levelname)s: %(message)s")
logger = logging.getLogger(__name__)

GROUND_TRUTH_PATH = Path("evaluation/ground_truth/gold_standard.xlsx")
RESULTS_MD_PATH = Path("docs/evaluation_results.md")
RAW_JSON_PATH = Path("results/baseline_comparison_raw.json")

BERT_MODEL = "roberta-large"

# Approximate per-million-token pricing (USD) at time of evaluation.
# Update these if the comparison is re-run later. Source: provider sites.
# We use a single blended rate per model (avg of input + output) for a
# rough cost estimate — the actual breakdown depends on prompt vs response
# token counts, which varies by paper.
MODEL_PRICING_USD_PER_M_TOKENS = {
    "gpt-4o-mini": 0.375,  # blended ~$0.15 input + $0.60 output / 2
    "meta-llama/llama-3.1-70b-instruct": 0.60,  # blended OpenRouter rate
}

# Conservative average: 8 fields x ~3000 tokens each (prompt + response).
# Used only when actual token counts aren't returned by the provider.
ESTIMATED_TOKENS_PER_PAPER = 24_000


@dataclass
class ModelResult:
    """Aggregate results for one model run.

    Attributes:
        model: Model identifier.
        bert_f1: Mean BERTScore F1 across all (paper, field) pairs.
        bert_precision: Mean BERTScore Precision.
        bert_recall: Mean BERTScore Recall.
        latency_seconds: Total wall-clock time for extraction.
        estimated_cost_usd: Approximate API cost based on token estimates.
        n_pairs: Number of (paper, field) pairs evaluated.
    """

    model: str
    bert_f1: float
    bert_precision: float
    bert_recall: float
    latency_seconds: float
    estimated_cost_usd: float
    n_pairs: int


def _load_ground_truth(path: Path) -> dict[str, dict[str, str]]:
    """Load gold-standard annotations from the Excel file.

    Args:
        path: Path to gold_standard.xlsx.

    Returns:
        Nested dict: {paper_name: {field_name: value}}.
    """
    wb = load_workbook(path, read_only=True)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    data: dict[str, dict[str, str]] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(headers, row))
        paper = record.get("paper_name")
        if not paper:
            continue
        data[paper] = {
            spec.name: (record.get(spec.name) or "").strip()
            for spec in FIELD_SPECS
        }
    return data


def _run_pipeline_for_model(
    paper_names: list[str],
    model: str,
) -> tuple[dict[str, dict[str, str]], float]:
    """Extract all 8 fields for each paper using the specified model.

    Args:
        paper_names: Filenames to extract.
        model: Model identifier (e.g., 'gpt-4o-mini' or
            'meta-llama/llama-3.1-70b-instruct').

    Returns:
        Tuple of (predictions_dict, latency_seconds).
    """
    store = VectorStore()
    indexed = set(store.list_papers())
    paper_names = [p for p in paper_names if p in indexed]
    searcher = HybridSearcher(store)

    predictions: dict[str, dict[str, str]] = {}
    start = time.time()
    for paper in paper_names:
        print(f"  [{model}] Extracting {paper} ...")
        ex = extract_paper(searcher, paper, model=model)
        predictions[paper] = {
            name: (fe.value or "") for name, fe in ex.fields.items()
        }
    latency = time.time() - start
    return predictions, latency


def _evaluate_predictions(
    predictions: dict[str, dict[str, str]],
    ground_truth: dict[str, dict[str, str]],
) -> tuple[float, float, float, int]:
    """Compute mean BERTScore P/R/F1 over all (paper, field) pairs.

    Args:
        predictions: Per-paper predictions {paper: {field: value}}.
        ground_truth: Per-paper references in same shape.

    Returns:
        Tuple (mean_precision, mean_recall, mean_f1, n_pairs).
    """
    pred_strings: list[str] = []
    ref_strings: list[str] = []
    for paper in ground_truth:
        if paper not in predictions:
            continue
        for spec in FIELD_SPECS:
            pred_strings.append(predictions[paper].get(spec.name, ""))
            ref_strings.append(ground_truth[paper].get(spec.name, ""))

    if not pred_strings:
        return 0.0, 0.0, 0.0, 0

    P, R, F1 = bert_score(
        cands=pred_strings,
        refs=ref_strings,
        model_type=BERT_MODEL,
        lang="en",
        verbose=False,
        rescale_with_baseline=False,
    )
    n = len(pred_strings)
    return (
        sum(P.tolist()) / n,
        sum(R.tolist()) / n,
        sum(F1.tolist()) / n,
        n,
    )


def _estimate_cost(model: str, n_papers: int) -> float:
    """Roughly estimate API cost for a model run over n papers.

    Token counts are estimated rather than measured because OpenRouter
    and OpenAI report usage slightly differently across SDK versions.
    The estimate is conservative; the README discusses these caveats.

    Args:
        model: Model identifier.
        n_papers: Number of papers extracted.

    Returns:
        Estimated cost in USD.
    """
    rate = MODEL_PRICING_USD_PER_M_TOKENS.get(model)
    if rate is None:
        return 0.0
    total_tokens = ESTIMATED_TOKENS_PER_PAPER * n_papers
    return (total_tokens / 1_000_000) * rate


def _append_to_markdown(report_text: str, output_path: Path) -> None:
    """Append/replace the baseline comparison section in evaluation_results.md."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    existing = output_path.read_text(encoding="utf-8") if output_path.exists() else ""
    marker = "## Baseline Comparison: GPT-4o-mini vs Llama-3.1-70b"
    if marker in existing:
        before = existing.split(marker)[0]
        rest = existing.split(marker, 1)[1]
        next_heading_idx = rest.find("\n## ")
        after = rest[next_heading_idx:] if next_heading_idx >= 0 else ""
        output_path.write_text(before + report_text + after, encoding="utf-8")
    else:
        output_path.write_text(existing + "\n" + report_text, encoding="utf-8")


def main() -> None:
    """Entry point: run pipeline with each model, score, compare, report."""
    print("Loading ground truth ...")
    ground_truth = _load_ground_truth(GROUND_TRUTH_PATH)
    paper_names = list(ground_truth.keys())
    print(f"  {len(paper_names)} papers, "
          f"{sum(len(f) for f in ground_truth.values())} field values")

    models_to_compare = [
        config.EXTRACTION_MODEL,  # gpt-4o-mini
        config.BASELINE_MODEL,    # meta-llama/llama-3.1-70b-instruct
    ]

    results: list[ModelResult] = []
    raw_predictions: dict[str, dict[str, dict[str, str]]] = {}

    for model in models_to_compare:
        print(f"\n--- Running model: {model} ---")
        predictions, latency = _run_pipeline_for_model(paper_names, model)
        raw_predictions[model] = predictions

        print(f"  [{model}] Computing BERTScore ...")
        p, r, f1, n = _evaluate_predictions(predictions, ground_truth)
        cost = _estimate_cost(model, len(paper_names))

        results.append(
            ModelResult(
                model=model,
                bert_f1=f1,
                bert_precision=p,
                bert_recall=r,
                latency_seconds=latency,
                estimated_cost_usd=cost,
                n_pairs=n,
            )
        )
        print(f"  [{model}] F1={f1:.4f}  latency={latency:.1f}s  "
              f"~${cost:.4f}")

    # Print side-by-side summary
    print("\n" + "=" * 70)
    print("BASELINE COMPARISON")
    print("=" * 70)
    print(f"{'Model':<40} {'F1':>8} {'Latency':>10} {'~Cost':>10}")
    print("-" * 70)
    for r in results:
        print(f"{r.model:<40} {r.bert_f1:>8.4f} "
              f"{r.latency_seconds:>9.1f}s {r.estimated_cost_usd:>9.4f}$")
    print("=" * 70)

    # Markdown report
    md_lines = [
        "## Baseline Comparison: GPT-4o-mini vs Llama-3.1-70b",
        "",
        f"**Evaluation set:** {len(paper_names)} papers, "
        f"{results[0].n_pairs if results else 0} (paper, field) pairs",
        "",
        "| Model | BERTScore F1 | BERTScore Precision | BERTScore Recall | Latency (s) | ~Cost (USD) |",
        "|---|---|---|---|---|---|",
    ]
    for r in results:
        md_lines.append(
            f"| `{r.model}` | **{r.bert_f1:.4f}** | {r.bert_precision:.4f} "
            f"| {r.bert_recall:.4f} | {r.latency_seconds:.1f} "
            f"| ~${r.estimated_cost_usd:.4f} |"
        )
    md_lines.extend(
        [
            "",
            "**Cost notes:** Cost is a rough estimate based on a fixed "
            "per-paper token average. Actual costs vary by paper length and "
            "extracted text length.",
            "",
            "**Latency notes:** Latency includes network round-trip plus "
            "model inference; OpenRouter routing through Llama-70b-instruct "
            "typically incurs higher latency than OpenAI's first-party API.",
            "",
        ]
    )
    _append_to_markdown("\n".join(md_lines), RESULTS_MD_PATH)

    # Persist raw outputs for inspection
    RAW_JSON_PATH.parent.mkdir(parents=True, exist_ok=True)
    with RAW_JSON_PATH.open("w", encoding="utf-8") as fh:
        json.dump(
            {
                "results": [
                    {
                        "model": r.model,
                        "bert_f1": r.bert_f1,
                        "bert_precision": r.bert_precision,
                        "bert_recall": r.bert_recall,
                        "latency_seconds": r.latency_seconds,
                        "estimated_cost_usd": r.estimated_cost_usd,
                        "n_pairs": r.n_pairs,
                    }
                    for r in results
                ],
                "raw_predictions": raw_predictions,
            },
            fh,
            indent=2,
            ensure_ascii=False,
        )

    print(f"\nAppended results to {RESULTS_MD_PATH}")
    print(f"Wrote {RAW_JSON_PATH}")


if __name__ == "__main__":
    main()